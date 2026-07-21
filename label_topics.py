#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_report.py が生成した「主な話題・テーマ」シートの「話題名」「概要」列に、
執筆済みのラベルを書き込む。

使い方:
    python3 label_topics.py labels.json

labels.json は debug/topics_for_labeling.json と同じ順番(件数の多い順)で、
各トピックに対して {"name": "話題名", "description": "1〜2文の概要"} を並べた
JSON配列。例:

[
  {"name": "行政手続き・許認可・例規改正",
   "description": "各種申請や契約書の条項確認、条例の新旧対照表作成など、制度運用に関する確認依頼が多い。"},
  ...
]

要素数は debug/topics_for_labeling.json のトピック数と一致させること
(一部だけ埋めたい場合は、埋めない要素を {} にしておけば "(要記入)" のまま残る)。
"""
import sys
import json
import glob
import os

HERE = os.path.dirname(os.path.abspath(__file__))


def label_workbook(path, labels):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    wb = load_workbook(path)
    if '主な話題・テーマ' not in wb.sheetnames:
        print(f"WARN: {path} に「主な話題・テーマ」シートがありません。スキップ。")
        return
    ws = wb['主な話題・テーマ']
    for i, label in enumerate(labels):
        row = 5 + i
        name = label.get('name')
        desc = label.get('description')
        if name:
            ws.cell(row, 2).value = name
        if desc:
            c = ws.cell(row, 6)
            c.value = desc
            c.alignment = Alignment(wrap_text=True, vertical='top')
    wb.save(path)
    print(f"更新: {path}")


def main():
    if len(sys.argv) < 2:
        print("使い方: python3 label_topics.py labels.json")
        sys.exit(1)
    with open(sys.argv[1], encoding='utf-8') as f:
        labels = json.load(f)

    month_files = sorted(glob.glob(os.path.join(HERE, 'reports', 'QommonsAI利用集計_*.xlsx')))
    if not month_files:
        print("ERROR: reports/QommonsAI利用集計_*.xlsx が見つかりません。先に build_report.py を実行してください。")
        sys.exit(1)
    label_workbook(month_files[-1], labels)
    latest = os.path.join(HERE, 'QommonsAI利用集計.xlsx')
    if os.path.exists(latest):
        label_workbook(latest, labels)


if __name__ == '__main__':
    main()

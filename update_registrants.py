#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
fetch_registrants.mjs がダウンロードした組織メンバー一覧(Excel)から
registrants.csv を再生成する。

使い方:
    python3 update_registrants.py [downloads/qommons-member-template-YYYY-MM-DD.xlsx]

引数を省略した場合は downloads/ 内の最新の qommons-member-template-*.xlsx を使う。

Excel「ユーザー」シートの構成(2026-08-07 時点):
    3行目  ヘッダー行
    4行目  各列の説明行(データではない)
    5行目〜 実データ(現在の登録メンバーが事前入力されている)

    C列 メールアドレス（ログインID）  … 種別=メールアドレスの場合
    D列 ユーザーID                    … 種別=ユーザーIDの場合(職員番号)
    F列 氏名 / G列 氏名(カナ)
    J列 アクセスレベル / K列 部署名
    O列 本パスワード設定状況(完了/未完了)  → registrants.csv の status

差分(追加/削除/部署変更/ステータス変更)を標準出力に表示してから書き出す。
"""
import sys
import os
import csv
import glob

HERE = os.path.dirname(os.path.abspath(__file__))
REGISTRANTS_PATH = os.path.join(HERE, 'registrants.csv')
EMPLOYEE_NO_LOGIN_SUFFIX = '@北海道_函館市'

FIELDS = ['email', 'employee_no', 'name', 'name_kana', 'access_level', 'department', 'status']

# Excel「ユーザー」シートの列番号(1始まり)
COL_EMAIL, COL_USER_ID, COL_NAME, COL_KANA = 3, 4, 6, 7
COL_ACCESS_LEVEL, COL_DEPARTMENT, COL_STATUS = 10, 11, 15
FIRST_DATA_ROW = 5


def login_key(rec):
    return rec['email'] if rec['email'] else f"{rec['employee_no']}{EMPLOYEE_NO_LOGIN_SUFFIX}"


def load_current():
    if not os.path.exists(REGISTRANTS_PATH):
        return {}
    out = {}
    with open(REGISTRANTS_PATH, encoding='utf-8-sig', newline='') as f:
        for row in csv.DictReader(f):
            rec = {k: (row.get(k) or '').strip() for k in FIELDS}
            out[login_key(rec)] = rec
    return out


def load_template(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True)['ユーザー']
    out = {}
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True):
        def cell(col):
            v = row[col - 1] if len(row) >= col else None
            return str(v).strip() if v is not None else ''
        name = cell(COL_NAME)
        email, employee_no = cell(COL_EMAIL), cell(COL_USER_ID)
        if not name or not (email or employee_no):
            continue  # 未使用のテンプレート行
        rec = {
            'email': email,
            'employee_no': employee_no,
            'name': name,
            'name_kana': cell(COL_KANA),
            'access_level': cell(COL_ACCESS_LEVEL),
            'department': cell(COL_DEPARTMENT),
            'status': cell(COL_STATUS),
        }
        out[login_key(rec)] = rec
    return out


def main():
    if len(sys.argv) > 1:
        path = sys.argv[1]
    else:
        files = sorted(glob.glob(os.path.join(HERE, 'downloads', 'qommons-member-template-*.xlsx')))
        if not files:
            raise SystemExit('ERROR: downloads/ に qommons-member-template-*.xlsx が見つかりません。'
                             '先に node fetch_registrants.mjs を実行してください。')
        path = files[-1]
    print(f'入力: {path}')

    current, fresh = load_current(), load_template(path)
    added = sorted(set(fresh) - set(current))
    removed = sorted(set(current) - set(fresh))
    dept_changed, status_changed = [], []
    for k in sorted(set(current) & set(fresh)):
        if current[k]['department'] != fresh[k]['department']:
            dept_changed.append((k, current[k]['department'], fresh[k]['department']))
        if current[k]['status'] != fresh[k]['status']:
            status_changed.append((k, current[k]['status'], fresh[k]['status']))

    print(f'登録者数: {len(current)} → {len(fresh)}')
    print(f'  追加: {len(added)} / 削除: {len(removed)} / '
          f'部署変更: {len(dept_changed)} / ステータス変更: {len(status_changed)}')
    for k, old, new in dept_changed:
        print(f'    部署変更: {k}: {old or "(空)"} → {new or "(空)"}')

    with open(REGISTRANTS_PATH, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=FIELDS)
        w.writeheader()
        for k in sorted(fresh):
            w.writerow(fresh[k])
    print(f'更新: {REGISTRANTS_PATH}')


if __name__ == '__main__':
    main()
